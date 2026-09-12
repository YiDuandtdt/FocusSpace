from pathlib import Path
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path('D:/19798/Download/Work/FocusSpace/Doc')
SOURCE = ROOT / 'FocusSpace_开发设计文档_v1.0.md'
OUT = ROOT / 'FocusSpace_开发设计文档_v1.0.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)

# compact_reference_guide; named overrides: ChineseFont, TableCompact,
# TechnicalMasthead (no separate cover), QuietFurniture.
def style(name, size, before=0, after=6, color='202B38', bold=False, spacing=1.25):
    s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.font.name = 'Calibri'
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.color.rgb = RGBColor.from_string(color)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
    f = s.paragraph_format
    f.space_before, f.space_after = Pt(before), Pt(after)
    f.line_spacing = spacing
    f.widow_control = True
    return s

style('Normal', 11)
style('Title', 26, after=9, color='0B2545', bold=True)
style('Subtitle', 10, after=12, color='667586')
for name, size, before, after, color in [('Heading 1',16,18,10,'2E74B5'), ('Heading 2',13,14,7,'2E74B5'), ('Heading 3',12,10,5,'1F4D78')]:
    s=style(name,size,before,after,color,True)
    s.paragraph_format.keep_with_next=True
style('Table Text',9.5,after=3,spacing=1.15)
style('Table Header',9.5,after=3,color='0B2545',bold=True,spacing=1.15)
style('Header',8.5,after=0,color='667586',spacing=1)
style('Footer',8.5,after=0,color='667586',spacing=1)

header=sec.header.paragraphs[0]
header.text='FOCUSSPACE  /  开发设计文档'
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('v1.0   ·   ')
field=OxmlElement('w:fldSimple')
field.set(qn('w:instr'),'PAGE')
footer._p.append(field)

def rich(p, text):
    for part in re.split(r'(\[[^\]]+\]\(https?://[^)]+\))',text):
        m=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',part)
        if not m:
            p.add_run(part)
            continue
        h=OxmlElement('w:hyperlink')
        h.set(qn('r:id'),p.part.relate_to(m.group(2),RT.HYPERLINK,is_external=True))
        r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
        col=OxmlElement('w:color'); col.set(qn('w:val'),'2E74B5'); pr.append(col)
        r.append(pr); t=OxmlElement('w:t'); t.text=m.group(1); r.append(t); h.append(r); p._p.append(h)

def make_table(rows):
    n=len(rows[0])
    # widths are explicit DXA and sum to the 9360-DXA content width.
    if n==2:
        widths=[2300,7060]
    elif n==3:
        widths=[2750,3700,2910]
    else:
        widths=[1700,2350,1450,3860]
    t=doc.add_table(rows=0, cols=n)
    t.autofit=False
    pr=t._tbl.tblPr
    pr.find(qn('w:tblW')).set(qn('w:type'),'dxa')
    pr.find(qn('w:tblW')).set(qn('w:w'),'9360')
    ind=OxmlElement('w:tblInd'); ind.set(qn('w:w'),'120'); ind.set(qn('w:type'),'dxa'); pr.append(ind)
    margins=OxmlElement('w:tblCellMar')
    for key,value in [('top',80),('bottom',80),('start',120),('end',120)]:
        x=OxmlElement('w:'+key); x.set(qn('w:w'),str(value)); x.set(qn('w:type'),'dxa'); margins.append(x)
    pr.append(margins)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+side); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D2DCE6'); borders.append(x)
    pr.append(borders)
    for c,w in zip(t._tbl.tblGrid,widths): c.set(qn('w:w'),str(w))
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        rp=cells[0]._tc.getparent().get_or_add_trPr()
        rp.append(OxmlElement('w:cantSplit'))
        if i==0: rp.append(OxmlElement('w:tblHeader'))
        for c,w,txt in zip(cells,widths,row):
            c.width=Inches(w/1440)
            p=c.paragraphs[0]; p.style='Table Header' if i==0 else 'Table Text'
            rich(p,txt)
            if i==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EEF5'); c._tc.get_or_add_tcPr().append(sh)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

lines=SOURCE.read_text(encoding='utf-8').splitlines()
i=0
while i<len(lines):
    line=lines[i].strip()
    if not line: i+=1; continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',x) for x in row): rows.append(row)
            i+=1
        make_table(rows)
        continue
    if line.startswith('# '):
        rich(doc.add_paragraph(style='Title'),line[2:])
    elif line.startswith('## '):
        rich(doc.add_paragraph(style='Heading 1'),line[3:])
    elif line.startswith('### '):
        rich(doc.add_paragraph(style='Heading 2'),line[4:])
    elif line.startswith('版本：'):
        rich(doc.add_paragraph(style='Subtitle'),line)
    else:
        rich(doc.add_paragraph(),line)
    i+=1
doc.core_properties.title='FocusSpace 开发设计文档 v1.0'
doc.core_properties.subject='MVP 架构、业务规则、数据模型与实现阶段'
doc.core_properties.author=''
doc.core_properties.keywords='FocusSpace,开发设计,MVP'
doc.save(OUT)

# Bounded structural check: file readability and complete table geometry.
check=Document(OUT)
for t in check.tables:
    widths=[int(c.get(qn('w:w'))) for c in t._tbl.tblGrid]
    assert sum(widths)==9360
    for row in t.rows:
        assert [int(c._tc.tcPr.tcW.w) for c in row.cells]==widths
assert len(check.tables)==9, len(check.tables)
print(f'Created: {OUT}\nParagraphs: {len(check.paragraphs)}; tables: {len(check.tables)}; source characters: {len(SOURCE.read_text(encoding="utf-8"))}')
